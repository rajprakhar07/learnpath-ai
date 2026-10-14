"""PPTX ingestion: one chunk per slide (split if very long), slide title as topic, pictures extracted as figures.

Notes:
  * Speaker notes are indexed as part of their slide (a citation to "slide 4" may quote that slide's notes).
  * Bullets without terminal punctuation get a "." appended so sentence-based features work; wording is unchanged.
  * Hidden slides are skipped (they are not shown to an audience). Slide numbers are the original deck positions.
  * Slide visuals are NOT rendered; the viewer shows the slide's text, notes and extracted pictures.
"""
import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError

from .ingest_pdf import IngestError, chunk_page, clean_heading

MIME = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif", "bmp": "image/bmp", "tiff": "image/tiff"}


def _terminate(t: str) -> str:
    t = re.sub(r"\s+", " ", t).strip()
    return t if not t or t[-1] in ".!?" else t + "."


def _walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(sh.shapes)
        else:
            yield sh


def read_slides(path):
    """-> list of dicts {slide, title, paragraphs, notes, figures}. Hidden slides omitted."""
    try:
        prs = Presentation(str(path))
    except (PackageNotFoundError, KeyError, ValueError, OSError) as e:
        raise IngestError("Could not read this PowerPoint file. Is it a valid .pptx (not .ppt)?") from e
    except Exception as e:  # python-pptx raises assorted errors on corrupt zips
        raise IngestError(f"Could not read this PowerPoint file ({type(e).__name__}).") from e
    slides = []
    for i, s in enumerate(prs.slides, start=1):
        if s._element.get("show") == "0":
            continue
        title_shape = s.shapes.title
        title = re.sub(r"\s+", " ", title_shape.text_frame.text).strip() if title_shape is not None and title_shape.has_text_frame else ""
        paras, figs = [], []
        for sh in _walk(s.shapes):
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img = sh.image
                    c = sh._element.xpath(".//p:cNvPr")
                    alt = (c[0].get("descr") or "").strip() if c else ""
                    ext = (img.ext or "png").lower()
                    w, h = img.size
                    figs.append({"blob": img.blob, "ext": ext, "mime": MIME.get(ext, "application/octet-stream"),
                                 "width": w, "height": h, "alt_text": alt, "caption": ""})
                except Exception:
                    pass
                continue
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = _terminate("".join(r.text for r in p.runs) or p.text)
                    if t:
                        paras.append(t)
            elif getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    cells = [re.sub(r"\s+", " ", c.text).strip() for c in row.cells]
                    t = _terminate(" | ".join(x for x in cells if x))
                    if t:
                        paras.append(t)
        notes = ""
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            notes = _terminate(s.notes_slide.notes_text_frame.text)
        slides.append({"slide": i, "title": title, "paragraphs": paras, "notes": notes, "figures": figs})
    return slides


def ingest_pptx(path):
    slides = read_slides(path)
    if not slides:
        raise IngestError("This presentation has no visible slides.")
    chunks, figures, topic, n_empty = [], [], "", 0
    for sl in slides:
        body = " ".join(sl["paragraphs"] + ([sl["notes"]] if sl["notes"] else []))
        text = (sl["title"] + "\n" if sl["title"] else "") + body
        if not (sl["title"] or body):
            n_empty += 1
        fallback = topic or f"Slide {sl['slide']}"
        parts, topic_end = chunk_page(text.strip(), fallback)
        topic = topic_end if sl["title"] else topic
        for t, tp in parts:
            chunks.append({"loc_kind": "slide", "loc_start": sl["slide"], "loc_end": sl["slide"], "text": t,
                           "topic": tp or f"Slide {sl['slide']}"})
        for f in sl["figures"]:
            figures.append(dict(f, loc_kind="slide", loc_start=sl["slide"]))
    if not chunks:
        raise IngestError("No text found in this presentation (it may contain only images).")
    return slides, chunks, figures, n_empty
