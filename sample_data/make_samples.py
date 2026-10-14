"""Generates the multi-format sample corpus. Run:  python sample_data/make_samples.py
  sample_course.pdf is made by make_sample_pdf.py; this script makes:
  sample_deck.pptx      databases deck: 6 slides, a picture with alt text, speaker notes, a table
  sample_lecture.vtt    operating-systems lecture transcript with timestamps
  sample_figure.pdf     2-page PDF with an embedded figure and a "Figure 1:" caption
"""
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image as RLImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

HERE = Path(__file__).parent


def make_png(path, kind):
    img = Image.new("RGB", (640, 360), "white")
    d = ImageDraw.Draw(img)
    if kind == "btree":
        d.rectangle([250, 30, 390, 80], outline="black", width=3)
        for i, x in enumerate((60, 260, 460)):
            d.rectangle([x, 220, x + 130, 270], outline="black", width=3)
            d.line([320, 80, x + 65, 220], fill="black", width=2)
    else:  # error curves
        d.line([60, 300, 600, 300], fill="black", width=2)
        d.line([60, 300, 60, 40], fill="black", width=2)
        d.line([60, 80, 200, 150, 340, 200, 600, 260], fill="blue", width=4)     # training error falls
        d.line([60, 90, 200, 160, 340, 190, 450, 170, 600, 100], fill="red", width=4)  # test error U-shape
    img.save(path)


def make_pptx(out):
    prs = Presentation()
    def slide(title, bullets, notes=None, picture=None, alt=None, table=None):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        body = s.placeholders[1].text_frame
        body.clear()
        for i, b in enumerate(bullets):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = b
            p.font.size = Pt(20)
        if picture:
            pic = s.shapes.add_picture(str(picture), Inches(5.2), Inches(3.4), width=Inches(4))
            pic._element.nvPicPr.cNvPr.set("descr", alt or "")
        if table:
            rows, cols = len(table), len(table[0])
            t = s.shapes.add_table(rows, cols, Inches(0.5), Inches(4.6), Inches(4.5), Inches(1.2)).table
            for r in range(rows):
                for c in range(cols):
                    t.cell(r, c).text = table[r][c]
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    t = prs.slides.add_slide(prs.slide_layouts[0])
    t.shapes.title.text = "Introduction to Relational Databases"
    t.placeholders[1].text = "CS240 sample deck"
    slide("Tables and Keys", [
        "A relational database stores data in tables made of rows and columns.",
        "A primary key uniquely identifies each row in a table.",
        "A foreign key is a column that references the primary key of another table."],
        table=[["Term", "Meaning"], ["Row", "One record"]])
    slide("SQL Basics", [
        "SELECT retrieves rows from one or more tables.",
        "The WHERE clause filters rows before they are returned.",
        "A JOIN combines rows from two tables based on a related column."])
    png = HERE / "_btree.png"
    make_png(png, "btree")
    slide("Indexes", [
        "An index is a data structure that speeds up lookups on a column.",
        "Indexes make reads faster but slow down writes, because the index must be updated."],
        notes="B-trees keep their keys sorted, so a lookup follows a short path from the root to a leaf.",
        picture=png, alt="Diagram of a B-tree index with a root node and three leaf nodes")
    png.unlink()
    slide("Normalization", [
        "Normalization organizes data to reduce redundancy.",
        "In third normal form, every non-key column depends only on the key."])
    slide("Transactions and ACID", [
        "A transaction is a group of operations that succeed or fail together.",
        "ACID stands for atomicity, consistency, isolation, and durability.",
        "Atomicity means that a transaction is applied completely or not at all."],
        notes="Durability means committed changes survive a crash.")
    prs.save(str(out))


VTT = """WEBVTT

00:00:03.000 --> 00:00:09.000
Welcome to today's lecture on operating systems.

00:00:20.000 --> 00:00:29.000
A process is a program in execution, with its own address space.

00:00:45.000 --> 00:00:56.000
A thread is a unit of execution inside a process, and threads in one process share memory.

00:01:30.000 --> 00:01:38.000
Scheduling decides which process runs on the CPU next.

00:02:10.000 --> 00:02:22.000
Round robin scheduling gives each process a fixed time slice called a quantum.

00:03:05.000 --> 00:03:17.000
If the quantum is too small, the system wastes time on context switches.

00:03:40.000 --> 00:03:50.000
A context switch saves the state of one process and loads the state of another.

00:05:12.000 --> 00:05:20.000
Deadlock occurs when processes wait for each other forever.

00:05:50.000 --> 00:06:08.000
The four necessary conditions for deadlock are mutual exclusion, hold and wait, no preemption, and circular wait.

00:07:15.000 --> 00:07:24.000
A semaphore is a counter used to control access to a shared resource.

00:08:02.000 --> 00:08:13.000
A mutex allows only one thread at a time to enter a critical section.
"""


def make_figure_pdf(out):
    png = HERE / "_curves.png"
    make_png(png, "curves")
    ss = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(out), pagesize=A4, leftMargin=2.4 * cm, rightMargin=2.4 * cm, title="Figure sample")
    story = [
        Paragraph("Model Complexity and Error", ss["Heading1"]),
        Paragraph("As model complexity grows, training error keeps falling, but test error first falls and then rises "
                  "again once the model begins to overfit the training data.", ss["BodyText"]),
        Spacer(1, 10), RLImage(str(png), width=11 * cm, height=6.2 * cm), Spacer(1, 6),
        Paragraph("Figure 1: Training error and test error versus model complexity.", ss["BodyText"]),
        PageBreak(),
        Paragraph("Choosing Complexity", ss["Heading1"]),
        Paragraph("The best model complexity is near the point where test error is lowest. "
                  "A validation set is used to find that point without touching the test set.", ss["BodyText"]),
    ]
    doc.build(story)
    png.unlink()


if __name__ == "__main__":
    make_pptx(HERE / "sample_deck.pptx")
    (HERE / "sample_lecture.vtt").write_text(VTT, encoding="utf-8")
    make_figure_pdf(HERE / "sample_figure.pdf")
    print("wrote sample_deck.pptx, sample_lecture.vtt, sample_figure.pdf")
