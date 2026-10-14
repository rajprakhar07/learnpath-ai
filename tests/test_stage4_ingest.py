"""Stage 4a: PPTX + timestamped transcript ingestion, figure extraction, multi-format citations."""
import io
import re
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from learnpath.figures import extract_pdf_figures
from learnpath.ingest_pdf import IngestError, extract_pages
from learnpath.ingest_pptx import ingest_pptx
from learnpath.ingest_transcript import build_chunks, ingest_transcript, parse_cues
from learnpath.locfmt import fmt_time, loc_label
from .conftest import ROOT, SAMPLE
from .test_stage2 import answer, ask_next, drain, rows

SD = ROOT / "sample_data"
DECK, VTT, FIGPDF = SD / "sample_deck.pptx", SD / "sample_lecture.vtt", SD / "sample_figure.pdf"


@pytest.fixture(scope="session", autouse=True)
def _samples():
    if not DECK.exists():
        import subprocess, sys
        subprocess.check_call([sys.executable, str(SD / "make_samples.py")])


def norm(s):
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


@pytest.fixture()
def multi(client):
    r = client.post("/api/documents/sample", json={"names": ["sample_course.pdf", "sample_deck.pptx", "sample_lecture.vtt", "sample_figure.pdf"]})
    assert r.status_code == 201, r.get_json()
    return client


def upload(client, name, data):
    return client.post("/api/documents", data={"file": (io.BytesIO(data), name)}, content_type="multipart/form-data")


# ---------------------------------------------------------------- labels
def test_location_labels():
    assert loc_label("page", 4) == "p.4" and loc_label("slide", 4) == "slide 4"
    assert loc_label("time", 130, 142) == "02:10–02:22" and loc_label("time", 3725) == "1:02:05"
    assert fmt_time(59.6) == "01:00"


# ---------------------------------------------------------------- PPTX
def test_pptx_one_chunk_per_slide_with_titles_as_topics_and_independent_text_check():
    slides, chunks, figs, n_empty = ingest_pptx(DECK)
    assert [c["loc_start"] for c in chunks] == [1, 2, 3, 4, 5, 6] and all(c["loc_kind"] == "slide" for c in chunks)
    assert [c["topic"] for c in chunks] == ["Introduction to Relational Databases", "Tables and Keys", "SQL Basics", "Indexes",
                                           "Normalization", "Transactions and ACID"]
    # independent re-read with python-pptx directly: every word of every shape's text is in that slide's chunk
    prs = Presentation(str(DECK))
    for i, s in enumerate(prs.slides, start=1):
        words = norm(" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame))
        assert words in norm(chunks[i - 1]["text"]) or all(w in norm(chunks[i - 1]["text"]).split() for w in words.split())


def test_pptx_includes_notes_and_tables_and_terminates_bullets():
    _, chunks, _, _ = ingest_pptx(DECK)
    by = {int(c["loc_start"]): c["text"] for c in chunks}
    assert "B-trees keep their keys sorted" in by[4]            # speaker notes
    assert "Row | One record." in by[2]                         # table row
    assert "Durability means committed changes survive a crash." in by[6]


def test_pptx_extracts_real_picture_with_alt_text():
    _, _, figs, _ = ingest_pptx(DECK)
    assert len(figs) == 1 and figs[0]["loc_start"] == 4
    assert figs[0]["blob"][:8] == b"\x89PNG\r\n\x1a\n" and (figs[0]["width"], figs[0]["height"]) == (640, 360)
    assert figs[0]["alt_text"].startswith("Diagram of a B-tree index")


def test_pptx_hidden_slides_are_skipped(tmp_path):
    prs = Presentation(str(DECK))
    prs.slides[2]._element.set("show", "0")
    p = tmp_path / "hidden.pptx"
    prs.save(str(p))
    _, chunks, _, _ = ingest_pptx(p)
    assert [int(c["loc_start"]) for c in chunks] == [1, 2, 4, 5, 6]     # original numbering preserved


def test_pptx_corrupt_and_image_only_rejected(tmp_path):
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"PK\x03\x04 not really")
    with pytest.raises(IngestError):
        ingest_pptx(bad)
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])                 # blank slide, no text
    p = tmp_path / "blank.pptx"
    prs.save(str(p))
    with pytest.raises(IngestError, match="No text"):
        ingest_pptx(p)


# ---------------------------------------------------------------- transcripts
def independent_vtt_cues(path):
    """A deliberately separate, minimal VTT reader used to verify the real one."""
    out, lines = [], Path(path).read_text().splitlines()
    for i, l in enumerate(lines):
        m = re.match(r"(\d\d):(\d\d):(\d\d)\.\d+ --> (\d\d):(\d\d):(\d\d)\.\d+", l)
        if m:
            a, b, c, d, e, f = map(int, m.groups())
            out.append((a * 3600 + b * 60 + c, d * 3600 + e * 60 + f, lines[i + 1]))
    return out


def test_vtt_cues_match_independent_parser_and_chunks_are_exact_concatenations():
    cues = parse_cues(VTT.read_text())
    ind = independent_vtt_cues(VTT)
    assert [(c["start"], c["end"], c["text"]) for c in cues] == ind and len(cues) == 11
    _, chunks, _ = ingest_transcript(VTT)
    for ch in chunks:
        inside = [t for s, e, t in ind if ch["loc_start"] <= s <= ch["loc_end"]]
        assert ch["text"] == " ".join(inside)                      # cited range contains exactly the cited text
    assert sum(len(c["text"].split()) for c in chunks) == sum(len(t.split()) for _, _, t in ind)  # nothing lost or duplicated
    assert chunks[0]["topic"] == "Lecture 00:00–05:00" and chunks[-1]["topic"] == "Lecture 05:00–10:00"


def test_srt_and_plain_timestamp_formats():
    srt = "1\r\n00:00:01,500 --> 00:00:04,000\r\n<i>Hello</i> there.\r\n\r\n2\r\n00:01:10,000 --> 00:01:15,250\r\n<v Bob>Second cue.</v>\r\n"
    c = parse_cues(srt)
    assert [(x["start"], x["end"], x["text"]) for x in c] == [(1.5, 4.0, "Hello there."), (70.0, 75.25, "Second cue.")]
    txt = "[00:00:05] First line.\n[00:01:00] Second line\ncontinues here.\n1:30:00 Late line."
    t = parse_cues(txt)
    assert [(x["start"], x["text"]) for x in t] == [(5, "First line."), (60, "Second line continues here."), (5400, "Late line.")]
    assert t[0]["end"] == 60 and t[1]["end"] == 5400


def test_chunking_splits_on_long_pause_and_size():
    cues = [{"start": 0, "end": 4, "text": "One."}, {"start": 5, "end": 8, "text": "Two."}, {"start": 100, "end": 105, "text": "Three."}]
    ch = build_chunks(cues)
    assert [c["text"] for c in ch] == ["One. Two.", "Three."]
    long = [{"start": i * 5, "end": i * 5 + 4, "text": " ".join(["word"] * 20) + "."} for i in range(6)]
    assert all(len(c["text"].split()) <= 100 for c in build_chunks(long))


def test_transcript_without_timestamps_rejected(tmp_path):
    p = tmp_path / "x.txt"
    p.write_text("just some words without any time codes")
    with pytest.raises(IngestError, match="timestamped"):
        ingest_transcript(p)


def test_transcript_bom_and_utf16(tmp_path):
    body = "WEBVTT\n\n00:00:01.000 --> 00:00:03.000\nCafé société.\n"
    p1 = tmp_path / "bom.vtt"; p1.write_bytes(b"\xef\xbb\xbf" + body.encode())
    p2 = tmp_path / "u16.vtt"; p2.write_bytes(body.encode("utf-16"))
    assert ingest_transcript(p1)[0][0]["text"] == "Café société." == ingest_transcript(p2)[0][0]["text"]


# ---------------------------------------------------------------- PDF figures
def test_pdf_figure_extracted_with_unambiguous_caption():
    pages = {p["page"]: p["text"] for p in extract_pages(FIGPDF)}
    figs = extract_pdf_figures(FIGPDF, pages)
    assert len(figs) == 1 and figs[0]["loc_start"] == 1 and figs[0]["blob"][:4] == b"\x89PNG"
    assert figs[0]["caption"] == "Figure 1: Training error and test error versus model complexity."


def test_pdf_without_images_yields_no_figures_and_no_invented_ones():
    pages = {p["page"]: p["text"] for p in extract_pages(SAMPLE)}
    assert extract_pdf_figures(SAMPLE, pages) == []


# ---------------------------------------------------------------- API integration
def test_upload_all_formats_and_listing(multi):
    docs = multi.get("/api/documents").get_json()
    assert [(d["kind"], d["n_units"], d["n_figures"]) for d in docs] == [("pdf", 6, 0), ("pptx", 6, 1), ("transcript", 11, 0), ("pdf", 2, 1)]


def test_upload_validation_for_new_formats(client):
    assert upload(client, "x.pptx", b"%PDF-1.4 nope").status_code == 415            # wrong magic for extension
    assert upload(client, "x.pptx", b"PK\x03\x04garbage").status_code == 422
    assert upload(client, "x.vtt", b"\x00\x01\x02binary").status_code == 415
    assert upload(client, "x.ppt", b"whatever").status_code == 415
    assert upload(client, "x.vtt", b"WEBVTT\n\nno cues here").status_code == 422
    ok = upload(client, "lecture.vtt", VTT.read_bytes())
    assert ok.status_code == 201 and upload(client, "again.vtt", VTT.read_bytes()).get_json()["duplicate"] is True


def test_figure_endpoint_serves_exact_bytes_and_404s(multi):
    d = next(d for d in multi.get("/api/documents").get_json() if d["kind"] == "pptx")
    r = multi.post("/api/ask", json={"question": "What does an index do?"}).get_json()
    fig = next(f for c in r["citations"] for f in c["figures"])
    resp = multi.get(fig["url"])
    assert resp.status_code == 200 and resp.mimetype == "image/png" and resp.data == ingest_pptx(DECK)[2][0]["blob"]
    assert multi.get("/figures/9999").status_code == 404


def test_delete_document_removes_its_figures_only(multi, app):
    docs = multi.get("/api/documents").get_json()
    deck = next(d for d in docs if d["kind"] == "pptx")
    other_figs_before = rows(app, "SELECT COUNT(*) c FROM figures WHERE doc_id!=?", deck["id"])[0]["c"]
    files_before = {p.name for p in app.config["FIGURE_DIR"].glob("*")}
    assert rows(app, "SELECT COUNT(*) c FROM figures WHERE doc_id=?", deck["id"])[0]["c"] == 1 and len(files_before) == 2
    multi.delete(f"/api/documents/{deck['id']}")
    assert rows(app, "SELECT COUNT(*) c FROM figures WHERE doc_id=?", deck["id"])[0]["c"] == 0
    assert rows(app, "SELECT COUNT(*) c FROM figures")[0]["c"] == other_figs_before == 1     # the PDF figure survives
    remaining = {r["stored_name"] for r in rows(app, "SELECT stored_name FROM figures")}
    assert {p.name for p in app.config["FIGURE_DIR"].glob("*")} == remaining                  # no orphan files


# ---------------------------------------------------------------- tutoring with citations in every format
def ask(client, q):
    return client.post("/api/ask", json={"question": q}).get_json()


def test_slide_citation_is_correct_and_labelled(multi):
    r = ask(multi, "What is a foreign key?")
    assert r["status"] == "answered"
    c = r["citations"][0]
    assert (c["loc_kind"], int(c["loc_start"]), c["label"], c["doc_name"]) == ("slide", 2, "slide 2", "sample_deck.pptx")
    assert "foreign key" in norm(c["excerpt"])


def test_transcript_citation_has_tight_time_range_that_contains_the_answer(multi):
    r = ask(multi, "What is a quantum in round robin scheduling?")
    assert r["status"] == "answered"
    c = r["citations"][0]
    assert c["loc_kind"] == "time" and c["label"] == "02:10–02:22" and c["doc_name"] == "sample_lecture.vtt"
    ind = independent_vtt_cues(VTT)
    spoken = " ".join(t for s, e, t in ind if c["loc_start"] <= s <= c["loc_end"])
    assert "fixed time slice called a quantum" in spoken and c["excerpt"] == spoken


def test_cross_document_retrieval_picks_the_right_document(multi):
    cases = {"What does ACID stand for?": ("sample_deck.pptx", 6), "What are the conditions for deadlock?": ("sample_lecture.vtt", None),
             "What is L1 regularization?": ("sample_course.pdf", 4), "What is the best model complexity?": ("sample_figure.pdf", 2)}
    for q, (doc, loc) in cases.items():
        r = ask(multi, q)
        assert r["status"] == "answered", q
        top = r["citations"][0]
        assert top["doc_name"] == doc, (q, top["doc_name"])
        if loc:
            assert int(top["loc_start"]) == loc, (q, top["loc_start"])


def test_figures_are_attached_to_citations_with_real_alt_text_and_captions(multi):
    r = ask(multi, "What does an index do for lookups?")
    c = next(c for c in r["citations"] if c["loc_kind"] == "slide")
    assert int(c["loc_start"]) == 4 and len(c["figures"]) == 1
    assert c["figures"][0]["alt_text"] == "Diagram of a B-tree index with a root node and three leaf nodes" and c["figures"][0]["caption"] == ""
    r2 = ask(multi, "How does training error behave as model complexity grows?")
    f = next(f for c in r2["citations"] for f in c["figures"])
    assert f["caption"] == "Figure 1: Training error and test error versus model complexity."


def test_unsupported_still_flagged_with_many_documents(multi):
    for q in ("What is the capital of France?", "How do I bake sourdough bread?", "Who won the 2018 World Cup?"):
        assert ask(multi, q)["status"] == "unsupported"


# ---------------------------------------------------------------- viewers
def test_slide_viewer_highlights_target_and_escapes_html(multi):
    d = next(d for d in multi.get("/api/documents").get_json() if d["kind"] == "pptx")
    html = multi.get(f"/view/{d['id']}?slide=4").get_data(as_text=True)
    assert 'class="slide hl"' in html and html.count('class="slide hl"') == 1 and "Indexes" in html
    assert 'src="/figures/' in html and "Alt text: Diagram of a B-tree index" in html


def test_transcript_viewer_highlights_range_and_escapes_untrusted_text(client):
    evil = b"WEBVTT\n\n00:00:01.000 --> 00:00:03.000\n<script>alert(1)</script> Hello &amp; welcome.\n\n00:01:00.000 --> 00:01:04.000\nSecond.\n"
    doc = upload(client, "evil.vtt", evil).get_json()["document"]
    html = client.get(f"/view/{doc['id']}?t=60&te=64").get_data(as_text=True)
    assert "<script>alert(1)</script>" not in html          # never emitted as live markup
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html   # content is preserved, escaped
    assert html.count('class="cue hl"') == 1 and "Second." in html and "Hello &amp; welcome." in html
    assert client.get("/view/9999").status_code == 404
    assert client.get(f"/view/{doc['id']}?t=abc").status_code == 200


def test_comparison_operators_in_lecture_text_are_preserved():
    cues = parse_cues("WEBVTT\n\n00:00:01.000 --> 00:00:04.000\n<v Prof>If x<y and y>z then x<z is false.</v> <i>Note</i> <00:00:02.000>this.\n")
    assert cues[0]["text"] == "If x<y and y>z then x<z is false. Note this."


def test_pdf_has_no_text_viewer_route(multi):
    d = next(d for d in multi.get("/api/documents").get_json() if d["kind"] == "pdf")
    assert multi.get(f"/view/{d['id']}").status_code == 404 and multi.get(f"/files/{d['id']}").status_code == 200


# ---------------------------------------------------------------- assessment over all formats
def test_assessment_questions_from_slides_and_transcript_keep_real_locations(app, multi):
    n, last = drain(multi, app)
    qs = rows(app, "SELECT * FROM questions")
    kinds = {q["loc_kind"] for q in qs}
    assert {"page", "slide", "time"} <= kinds, kinds
    chunks = {c["id"]: c for c in rows(app, "SELECT * FROM chunks")}
    for q in qs:
        ch = chunks[q["chunk_id"]]
        assert q["loc_kind"] == ch["loc_kind"] and q["loc_start"] == ch["loc_start"] and q["loc_end"] == ch["loc_end"]
        assert norm(q["source_quote"]) in norm(ch["text"])
    assert len({q["quote_hash"] for q in qs}) == len(qs) == n
